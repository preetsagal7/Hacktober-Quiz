"""
Excel Exporter Service
Generates comprehensive multi-sheet Excel reports containing:
- Master Participants List with full personal & academic details + round breakdown
- Overall Winners & Leaderboard across all tournament rounds
- Individual Round Sheets with rankings, scores, and qualification/winner status
"""
import io
from datetime import datetime
from typing import List, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_dt(dt: Optional[datetime]) -> str:
    if not dt:
        return "-"
    try:
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(dt)


def format_duration(start: Optional[datetime], end: Optional[datetime]) -> str:
    if not start or not end:
        return "-"
    try:
        seconds = int((end - start).total_seconds())
        if seconds < 0:
            return "-"
        mins = seconds // 60
        secs = seconds % 60
        return f"{mins}m {secs}s"
    except Exception:
        return "-"


def style_table_headers(ws, row_idx: int, col_count: int, bg_hex: str = "1E293B"):
    header_fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='475569'),
        right=Side(style='thin', color='475569'),
        top=Side(style='thin', color='475569'),
        bottom=Side(style='thin', color='475569'),
    )
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[row_idx].height = 28


def autofit_columns(ws, min_width: int = 12, max_width: int = 45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title merged banner row lengths
            if cell.row in (1, 2):
                continue
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        adjusted_width = max(min_width, min(max_len + 4, max_width))
        ws.column_dimensions[col_letter].width = adjusted_width


def generate_participants_and_winners_excel(
    event: Any,
    participants: List[Any],
    rounds: List[Any],
    attempts: List[Any],
    qualifications: List[Any],
) -> io.BytesIO:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Pre-index attempts and qualifications for O(1) lookups
    attempts_map = {}
    for att in attempts:
        attempts_map[(att.participant_id, att.round_id)] = att

    quals_map = {}
    for q in qualifications:
        quals_map[(q.participant_id, q.round_id)] = q

    # Common styles
    font_family = "Segoe UI"
    data_font = Font(name=font_family, size=9)
    bold_font = Font(name=font_family, size=9, bold=True)
    mono_font = Font(name="Consolas", size=9, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    gold_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    silver_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bronze_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    event_name = getattr(event, "name", "CyberQuiz 2026")
    college_name = getattr(event, "college_name", "CyberQuiz Event") or "CyberQuiz Event"

    # Compute overall participant statistics
    participant_stats = []
    for p in participants:
        total_score = 0.0
        total_correct = 0
        total_incorrect = 0
        total_unanswered = 0
        rounds_participated = 0
        round_scores = {}

        for r in rounds:
            att = attempts_map.get((p.id, r.id))
            q = quals_map.get((p.id, r.id))
            if att:
                score = att.score or 0.0
                total_score += score
                total_correct += (att.correct_count or 0)
                total_incorrect += (att.incorrect_count or 0)
                total_unanswered += (att.unanswered_count or 0)
                rounds_participated += 1
                round_scores[r.id] = {
                    "score": score,
                    "status": att.status.value if hasattr(att.status, "value") else str(att.status),
                    "correct": att.correct_count or 0,
                    "incorrect": att.incorrect_count or 0,
                    "unanswered": att.unanswered_count or 0,
                    "qual": q,
                }
            elif q:
                total_score += (q.score or 0.0)
                round_scores[r.id] = {
                    "score": q.score or 0.0,
                    "status": "QUALIFIED" if q.is_qualified else "COMPLETED",
                    "correct": 0,
                    "incorrect": 0,
                    "unanswered": 0,
                    "qual": q,
                }
            else:
                round_scores[r.id] = None

        participant_stats.append({
            "participant": p,
            "total_score": round(total_score, 2),
            "total_correct": total_correct,
            "total_incorrect": total_incorrect,
            "total_unanswered": total_unanswered,
            "rounds_participated": rounds_participated,
            "round_scores": round_scores,
        })

    # Sort participants for Overall Winners (active participants first, then by total score desc, correct desc)
    overall_ranked = sorted(
        participant_stats,
        key=lambda item: (
            1 if item["rounds_participated"] > 0 else 0,
            item["total_score"],
            item["total_correct"],
        ),
        reverse=True
    )

    # ─────────────────────────────────────────────────────────────
    # SHEET 1: OVERALL WINNERS & TOURNAMENT STANDINGS
    # ─────────────────────────────────────────────────────────────
    ws_winners = wb.create_sheet(title="Overall Winners & Standings")
    ws_winners.views.sheetView[0].showGridLines = True

    # Title & Banner
    ws_winners.merge_cells("A1:K1")
    t1 = ws_winners["A1"]
    t1.value = f"{event_name.upper()} — OVERALL TOURNAMENT WINNERS & STANDINGS"
    t1.font = Font(name=font_family, size=13, bold=True, color="1E293B")
    t1.alignment = Alignment(horizontal="left", vertical="center")
    ws_winners.row_dimensions[1].height = 24

    ws_winners.merge_cells("A2:K2")
    t2 = ws_winners["A2"]
    t2.value = f"College / Organization: {college_name} | Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Enrolled: {len(participants)}"
    t2.font = Font(name=font_family, size=9, italic=True, color="64748B")
    t2.alignment = Alignment(horizontal="left", vertical="center")
    ws_winners.row_dimensions[2].height = 18

    # Table Headers for Overall Sheet
    w_headers = [
        "Overall Rank", "Standing / Award", "USN", "Participant Name", "College",
        "Department", "Semester", "Total Score", "Total Correct", "Rounds Completed"
    ]
    for r in rounds:
        w_headers.append(f"R{r.round_number}: {r.name[:18]} (Score)")
        w_headers.append(f"R{r.round_number} Result")

    header_row_w = 4
    for c_idx, h in enumerate(w_headers, start=1):
        cell = ws_winners.cell(row=header_row_w, column=c_idx, value=h)
    style_table_headers(ws_winners, header_row_w, len(w_headers), bg_hex="1E3A8A")  # Deep Navy Blue

    # Populate Overall Data
    cur_row = header_row_w + 1
    rank_counter = 1

    for item in overall_ranked:
        p = item["participant"]
        tot_sc = item["total_score"]
        tot_corr = item["total_correct"]
        participated = item["rounds_participated"]

        if participated > 0:
            rank_str = str(rank_counter)
            if rank_counter == 1:
                award = "🥇 1st Place - Champion"
                row_fill = gold_fill
            elif rank_counter == 2:
                award = "🥈 2nd Place - 1st Runner Up"
                row_fill = silver_fill
            elif rank_counter == 3:
                award = "🥉 3rd Place - 2nd Runner Up"
                row_fill = bronze_fill
            elif rank_counter <= 10:
                award = "🎖️ Top 10 Finalist"
                row_fill = green_fill
            else:
                award = "Participant"
                row_fill = zebra_fill if (rank_counter % 2 == 0) else None
            rank_counter += 1
        else:
            rank_str = "-"
            award = "Did Not Attempt"
            row_fill = None

        row_vals = [
            rank_str,
            award,
            p.usn,
            p.name or "Registered Student",
            p.college_name or college_name,
            p.department or "-",
            p.semester or "-",
            tot_sc,
            tot_corr,
            participated,
        ]

        for r in rounds:
            r_info = item["round_scores"].get(r.id)
            if r_info:
                row_vals.append(r_info["score"])
                q = r_info["qual"]
                if q:
                    res_str = f"Rank {q.rank} (Qualified)" if q.is_qualified else f"Rank {q.rank}"
                else:
                    res_str = r_info["status"]
                row_vals.append(res_str)
            else:
                row_vals.append("-")
                row_vals.append("Not Attempted")

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_winners.cell(row=cur_row, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if c_idx in (1, 8, 9, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                cell.font = mono_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws_winners.row_dimensions[cur_row].height = 20
        cur_row += 1

    ws_winners.freeze_panes = f"A{header_row_w + 1}"
    autofit_columns(ws_winners)

    # ─────────────────────────────────────────────────────────────
    # SHEET 2: ALL PARTICIPANTS MASTER LIST
    # ─────────────────────────────────────────────────────────────
    ws_master = wb.create_sheet(title="All Participants Master")
    ws_master.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_master.merge_cells("A1:M1")
    m1 = ws_master["A1"]
    m1.value = f"{event_name.upper()} — COMPLETE PARTICIPANT MASTER DIRECTORY"
    m1.font = Font(name=font_family, size=13, bold=True, color="1E293B")
    m1.alignment = Alignment(horizontal="left", vertical="center")
    ws_master.row_dimensions[1].height = 24

    ws_master.merge_cells("A2:M2")
    m2 = ws_master["A2"]
    m2.value = f"Total Registered: {len(participants)} Participants | College: {college_name} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    m2.font = Font(name=font_family, size=9, italic=True, color="64748B")
    m2.alignment = Alignment(horizontal="left", vertical="center")
    ws_master.row_dimensions[2].height = 18

    master_headers = [
        "S.No", "USN", "Full Name", "Roll Number", "Email Address",
        "Department", "College Name", "Semester", "Eligibility",
        "Ineligibility Reason", "Verified", "Registered Date", "Last Login",
        "Total Score", "Rounds Taken"
    ]
    for r in rounds:
        master_headers.append(f"R{r.round_number} Score")
        master_headers.append(f"R{r.round_number} Status")
        master_headers.append(f"R{r.round_number} Correct")
        master_headers.append(f"R{r.round_number} Standing")

    header_row_m = 4
    for c_idx, h in enumerate(master_headers, start=1):
        ws_master.cell(row=header_row_m, column=c_idx, value=h)
    style_table_headers(ws_master, header_row_m, len(master_headers), bg_hex="0F172A")

    # Sort participants alphabetically by USN
    sorted_by_usn = sorted(participant_stats, key=lambda x: x["participant"].usn)
    cur_row = header_row_m + 1

    for idx, item in enumerate(sorted_by_usn, start=1):
        p = item["participant"]
        tot_sc = item["total_score"]
        participated = item["rounds_participated"]

        is_eligible = getattr(p, "is_eligible", True)
        elig_str = "Eligible" if is_eligible else "Revoked / Ineligible"

        row_vals = [
            idx,
            p.usn,
            p.name or "Registered Student",
            p.roll_number or "-",
            p.email or "-",
            p.department or "-",
            p.college_name or college_name,
            p.semester or "-",
            elig_str,
            p.ineligibility_reason or "-",
            "Yes" if getattr(p, "is_verified", False) else "No",
            format_dt(p.created_at),
            format_dt(p.last_login_at),
            tot_sc,
            participated,
        ]

        for r in rounds:
            r_info = item["round_scores"].get(r.id)
            if r_info:
                row_vals.append(r_info["score"])
                row_vals.append(r_info["status"])
                row_vals.append(r_info["correct"])
                q = r_info["qual"]
                if q:
                    std = f"Rank {q.rank} (Qualified)" if q.is_qualified else f"Rank {q.rank}"
                else:
                    std = "Recorded"
                row_vals.append(std)
            else:
                row_vals.append("-")
                row_vals.append("Not Attempted")
                row_vals.append("-")
                row_vals.append("-")

        row_fill = None
        if not is_eligible:
            row_fill = red_fill
        elif idx % 2 == 0:
            row_fill = zebra_fill

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_master.cell(row=cur_row, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if c_idx in (1, 8, 9, 11, 14, 15):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:
                cell.font = mono_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws_master.row_dimensions[cur_row].height = 20
        cur_row += 1

    ws_master.freeze_panes = f"A{header_row_m + 1}"
    autofit_columns(ws_master)

    # ─────────────────────────────────────────────────────────────
    # SHEETS 3..N: DEDICATED SHEET FOR EACH ROUND (WINNERS & SCORES)
    # ─────────────────────────────────────────────────────────────
    participants_by_id = {p.id: p for p in participants}

    for r in rounds:
        safe_sheet_title = f"Round {r.round_number} - Winners"[:31]
        ws_r = wb.create_sheet(title=safe_sheet_title)
        ws_r.views.sheetView[0].showGridLines = True

        # Sheet Banner
        ws_r.merge_cells("A1:K1")
        r1 = ws_r["A1"]
        r1.value = f"ROUND {r.round_number}: {r.name.upper()} — WINNERS & LEADERBOARD"
        r1.font = Font(name=font_family, size=13, bold=True, color="1E293B")
        r1.alignment = Alignment(horizontal="left", vertical="center")
        ws_r.row_dimensions[1].height = 24

        qual_info = ""
        if getattr(r, "qualification_count", None):
            qual_info = f"Top {r.qualification_count} Qualify"
        elif getattr(r, "qualification_percentage", None):
            qual_info = f"Top {r.qualification_percentage}% Qualify"
        else:
            qual_info = "All Valid Submissions"

        ws_r.merge_cells("A2:K2")
        r2 = ws_r["A2"]
        r2.value = f"Status: {r.status.value.upper() if hasattr(r.status, 'value') else str(r.status).upper()} | Max Points: {r.max_score} | Duration: {r.duration_minutes} Mins | Questions: {r.question_count} | Criteria: {qual_info}"
        r2.font = Font(name=font_family, size=9, italic=True, color="64748B")
        r2.alignment = Alignment(horizontal="left", vertical="center")
        ws_r.row_dimensions[2].height = 18

        r_headers = [
            "Rank", "Standing / Status", "USN", "Participant Name", "Roll Number",
            "Email", "Department", "College", "Score", "Max Points", "Correct",
            "Incorrect", "Unanswered", "Exam Status", "Started At", "Submitted At", "Time Taken"
        ]
        header_row_r = 4
        for c_idx, h in enumerate(r_headers, start=1):
            ws_r.cell(row=header_row_r, column=c_idx, value=h)
        style_table_headers(ws_r, header_row_r, len(r_headers), bg_hex="065F46")  # Deep Emerald Green

        # Collect round attempts and qualifications
        r_attempts = [a for a in attempts if a.round_id == r.id]
        r_quals = [q for q in qualifications if q.round_id == r.id]

        # Map qualifications by participant
        r_qual_by_p = {q.participant_id: q for q in r_quals}

        # Build list of participants who took this round
        round_entries = []
        for att in r_attempts:
            p = participants_by_id.get(att.participant_id)
            if not p:
                continue
            q = r_qual_by_p.get(p.id)
            score = att.score if att.score is not None else (q.score if q else 0.0)
            round_entries.append({
                "participant": p,
                "attempt": att,
                "qualification": q,
                "score": score,
                "correct": att.correct_count or 0,
                "incorrect": att.incorrect_count or 0,
                "unanswered": att.unanswered_count or 0,
                "submitted_at": att.submitted_at,
                "started_at": att.started_at,
                "status": att.status.value if hasattr(att.status, "value") else str(att.status),
            })

        # Add any qualifications that might not have an attempt row
        for q in r_quals:
            if not any(e["participant"].id == q.participant_id for e in round_entries):
                p = participants_by_id.get(q.participant_id)
                if p:
                    round_entries.append({
                        "participant": p,
                        "attempt": None,
                        "qualification": q,
                        "score": q.score or 0.0,
                        "correct": 0,
                        "incorrect": 0,
                        "unanswered": 0,
                        "submitted_at": None,
                        "started_at": None,
                        "status": "QUALIFIED" if q.is_qualified else "RECORDED",
                    })

        # Sort round entries: official rank if qualification exists, otherwise score desc, submitted_at asc
        def sort_key(e):
            q = e["qualification"]
            if q and q.rank is not None:
                return (0, q.rank)
            return (1, -e["score"], e["submitted_at"] or datetime.max)

        sorted_round = sorted(round_entries, key=sort_key)

        # Qualification threshold calculation if not officially computed
        cutoff_count = len(sorted_round)
        if getattr(r, "qualification_count", None):
            cutoff_count = r.qualification_count
        elif getattr(r, "qualification_percentage", None) and len(sorted_round) > 0:
            cutoff_count = max(1, int(len(sorted_round) * r.qualification_percentage / 100))

        cur_row = header_row_r + 1
        for idx, entry in enumerate(sorted_round, start=1):
            p = entry["participant"]
            att = entry["attempt"]
            q = entry["qualification"]

            rank_val = q.rank if (q and q.rank) else idx
            is_qual = q.is_qualified if q else (idx <= cutoff_count)

            if rank_val == 1:
                standing = "🏆 1st Place Winner"
                row_fill = gold_fill
            elif rank_val == 2:
                standing = "🥈 2nd Place Winner"
                row_fill = silver_fill
            elif rank_val == 3:
                standing = "🥉 3rd Place Winner"
                row_fill = bronze_fill
            elif is_qual:
                standing = "✅ QUALIFIED FOR NEXT ROUND"
                row_fill = green_fill
            else:
                standing = "Completed"
                row_fill = zebra_fill if (idx % 2 == 0) else None

            duration_str = format_duration(entry["started_at"], entry["submitted_at"])

            row_vals = [
                rank_val,
                standing,
                p.usn,
                p.name or "Student",
                p.roll_number or "-",
                p.email or "-",
                p.department or "-",
                p.college_name or college_name,
                entry["score"],
                r.max_score,
                entry["correct"],
                entry["incorrect"],
                entry["unanswered"],
                entry["status"],
                format_dt(entry["started_at"]),
                format_dt(entry["submitted_at"]),
                duration_str,
            ]

            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_r.cell(row=cur_row, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                if row_fill:
                    cell.fill = row_fill
                if c_idx in (1, 9, 10, 11, 12, 13, 14, 17):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 3:
                    cell.font = mono_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws_r.row_dimensions[cur_row].height = 20
            cur_row += 1

        ws_r.freeze_panes = f"A{header_row_r + 1}"
        autofit_columns(ws_r)

    # Save to BytesIO buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
